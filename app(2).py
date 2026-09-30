import io
import json
import os
import re
from pathlib import Path
from typing import Any, Dict, List

import streamlit as st
from docx import Document
from groq import Groq
from pypdf import PdfReader

from prompts import build_resume_analysis_prompt, OUTPUT_SCHEMA

st.set_page_config(
    page_title="ATS Resume Analyzer AI",
    page_icon="📄",
    layout="wide",
)

APP_TITLE = "ATS Resume Analyzer AI"

SCORE_LIMITS = {
    "ats_formatting": 15,
    "keyword_match": 25,
    "professional_summary": 10,
    "work_experience": 20,
    "skills": 10,
    "education_certifications": 5,
    "projects_achievements": 10,
    "content_quality": 5,
}

DEFAULT_MODEL = os.getenv("GROQ_MODEL", "")

# -----------------------------
# File extraction
# -----------------------------

def clean_text(text: str) -> str:
    text = text.replace("\x00", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def extract_pdf(file_bytes: bytes) -> str:
    reader = PdfReader(io.BytesIO(file_bytes))
    pages = [(page.extract_text() or "") for page in reader.pages]
    return clean_text("\n\n".join(pages))


def extract_docx(file_bytes: bytes) -> str:
    document = Document(io.BytesIO(file_bytes))
    parts = [p.text for p in document.paragraphs if p.text.strip()]

    # Include table text because some resumes use tables.
    for table in document.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells)
            if row_text.strip():
                parts.append(row_text)

    return clean_text("\n".join(parts))


def extract_txt(file_bytes: bytes) -> str:
    return clean_text(file_bytes.decode("utf-8", errors="replace"))


def extract_uploaded_file(uploaded_file) -> str:
    if uploaded_file is None:
        raise ValueError("No file was uploaded.")

    suffix = Path(uploaded_file.name).suffix.lower()
    data = uploaded_file.getvalue()

    if suffix == ".pdf":
        text = extract_pdf(data)
    elif suffix == ".docx":
        text = extract_docx(data)
    elif suffix == ".txt":
        text = extract_txt(data)
    else:
        raise ValueError("Unsupported file type. Use PDF, DOCX, or TXT.")

    if not text.strip():
        raise ValueError(
            "No readable text was extracted. If this is a scanned/image-only "
            "PDF, OCR is required before analysis."
        )

    return text


# -----------------------------
# AI response validation
# -----------------------------

def clamp_score(value: Any, maximum: int) -> int:
    try:
        number = int(round(float(value)))
    except (TypeError, ValueError):
        number = 0
    return max(0, min(maximum, number))


def status_from_score(score: int) -> str:
    if score >= 85:
        return "Excellent ATS Compatibility"
    if score >= 70:
        return "Good"
    if score >= 55:
        return "Needs Improvement"
    return "Poor ATS Compatibility"


def default_result() -> Dict[str, Any]:
    return {
        "overall_score": 0,
        "status": "Poor ATS Compatibility",
        "summary": "",
        "category_scores": {key: 0 for key in SCORE_LIMITS},
        "strengths": [],
        "weaknesses": [],
        "matched_keywords": [],
        "missing_keywords": [],
        "unverified_keywords": [],
        "technical_skills_found": [],
        "technical_skills_missing": [],
        "soft_skills_found": [],
        "soft_skills_missing": [],
        "formatting_issues": [],
        "summary_feedback": "",
        "improved_summary": "",
        "experience_feedback": [],
        "project_feedback": [],
        "education_feedback": [],
        "achievement_feedback": [],
        "grammar_feedback": [],
        "bullet_improvements": [],
        "recommendations": [],
        "final_checklist": [],
    }


def normalize_result(data: Dict[str, Any]) -> Dict[str, Any]:
    result = default_result()

    if not isinstance(data, dict):
        data = {}

    for field in result:
        if field != "category_scores" and field in data:
            result[field] = data[field]

    incoming_scores = data.get("category_scores", {})
    if not isinstance(incoming_scores, dict):
        incoming_scores = {}

    total = 0
    for key, maximum in SCORE_LIMITS.items():
        score = clamp_score(incoming_scores.get(key, 0), maximum)
        result["category_scores"][key] = score
        total += score

    result["overall_score"] = total
    result["status"] = status_from_score(total)

    list_fields = [
        "strengths",
        "weaknesses",
        "matched_keywords",
        "missing_keywords",
        "unverified_keywords",
        "technical_skills_found",
        "technical_skills_missing",
        "soft_skills_found",
        "soft_skills_missing",
        "formatting_issues",
        "experience_feedback",
        "project_feedback",
        "education_feedback",
        "achievement_feedback",
        "grammar_feedback",
        "bullet_improvements",
        "recommendations",
        "final_checklist",
    ]

    for field in list_fields:
        value = result.get(field)
        if not isinstance(value, list):
            result[field] = [str(value)] if value else []

    for field in ["summary", "summary_feedback", "improved_summary"]:
        if not isinstance(result[field], str):
            result[field] = str(result[field] or "")

    return result


def parse_json_response(content: str) -> Dict[str, Any]:
    text = (content or "").strip()

    # Remove accidental Markdown code fences.
    text = re.sub(r"^```(?:json)?\s*", "", text, flags=re.IGNORECASE)
    text = re.sub(r"\s*```$", "", text).strip()

    try:
        return json.loads(text)
    except json.JSONDecodeError:
        start = text.find("{")
        end = text.rfind("}")
        if start >= 0 and end > start:
            return json.loads(text[start:end + 1])
        raise


# -----------------------------
# Groq analysis
# -----------------------------

def get_groq_credentials(model_override: str = ""):
    api_key = os.getenv("GROQ_API_KEY", "").strip()

    if not api_key:
        try:
            api_key = st.secrets.get("GROQ_API_KEY", "")
        except Exception:
            api_key = ""

    model = (model_override or os.getenv("GROQ_MODEL", "")).strip()

    if not model:
        try:
            model = st.secrets.get("GROQ_MODEL", "")
        except Exception:
            model = ""

    return api_key, model


def analyze_with_groq(
    resume_text: str,
    job_description: str,
    job_title: str,
    years_experience: str,
    model_override: str = "",
) -> Dict[str, Any]:

    api_key, model = get_groq_credentials(model_override)

    if not api_key:
        raise RuntimeError(
            "GROQ_API_KEY is missing. Add it in Streamlit Secrets or as an "
            "environment variable."
        )

    if not model:
        raise RuntimeError(
            "GROQ_MODEL is missing. Set it to a Groq model currently available "
            "to your account."
        )

    resume_text = resume_text[:45000]
    job_description = job_description[:30000]

    user_prompt = f"""
TARGET JOB TITLE:
{job_title or "Not provided"}

YEARS OF EXPERIENCE:
{years_experience or "Not provided"}

JOB DESCRIPTION:
{job_description or "Not provided. Perform resume-quality analysis without job-specific keyword matching."}

RESUME TEXT:
{resume_text}
"""

    client = Groq(api_key=api_key)

    try:
        response = client.chat.completions.create(
            model=model,
            temperature=0.1,
            max_tokens=5000,
            messages=[
                {
                    "role": "system",
                    "content": build_resume_analysis_prompt(),
                },
                {
                    "role": "user",
                    "content": user_prompt,
                },
            ],
        )
    except Exception as exc:
        message = str(exc)

        if "model" in message.lower() and (
            "not found" in message.lower()
            or "does not exist" in message.lower()
            or "access" in message.lower()
        ):
            raise RuntimeError(
                f"Groq model error: '{model}' is unavailable or inaccessible. "
                "Change GROQ_MODEL to a model currently available in your Groq account."
            ) from exc

        raise RuntimeError(f"Groq API request failed: {message}") from exc

    content = response.choices[0].message.content or ""

    try:
        data = parse_json_response(content)
    except Exception as exc:
        raise RuntimeError(
            "The AI returned invalid JSON. Please retry the analysis. "
            f"Response preview: {content[:400]}"
        ) from exc

    return normalize_result(data)


# -----------------------------
# Report generation
# -----------------------------

def markdown_list(items: List[Any]) -> str:
    if not items:
        return "_None identified._"
    return "\n".join(f"- {item}" for item in items)


def make_report(result: Dict[str, Any], filename: str) -> str:
    scores = result["category_scores"]

    score_lines = [
        "| Category | Score | Maximum |",
        "|---|---:|---:|",
    ]

    labels = {
        "ats_formatting": "ATS Formatting",
        "keyword_match": "Keyword Match",
        "professional_summary": "Professional Summary",
        "work_experience": "Work Experience",
        "skills": "Skills",
        "education_certifications": "Education & Certifications",
        "projects_achievements": "Projects & Achievements",
        "content_quality": "Content Quality",
    }

    for key, label in labels.items():
        score_lines.append(
            f"| {label} | {scores[key]} | {SCORE_LIMITS[key]} |"
        )

    return f"""# ATS Resume Analysis

**Resume:** {filename}

## Overall Result

**ATS Compatibility Score:** {result["overall_score"]}/100

**Status:** {result["status"]}

> This score is an estimated resume-to-job compatibility score. It is not a prediction of interview or hiring success.

## Summary

{result["summary"]}

## Category Scores

{chr(10).join(score_lines)}

## Strengths

{markdown_list(result["strengths"])}

## Weaknesses

{markdown_list(result["weaknesses"])}

## Matched Keywords

{markdown_list(result["matched_keywords"])}

## Missing Keywords

{markdown_list(result["missing_keywords"])}

## Unverified Keywords

{markdown_list(result["unverified_keywords"])}

## Technical Skills Found

{markdown_list(result["technical_skills_found"])}

## Technical Skills Missing

{markdown_list(result["technical_skills_missing"])}

## Soft Skills Found

{markdown_list(result["soft_skills_found"])}

## Soft Skills Missing

{markdown_list(result["soft_skills_missing"])}

## Formatting Issues

{markdown_list(result["formatting_issues"])}

## Professional Summary Feedback

{result["summary_feedback"]}

### Suggested Summary

{result["improved_summary"] or "_No suggested summary generated._"}

## Work Experience Feedback

{markdown_list(result["experience_feedback"])}

## Project Feedback

{markdown_list(result["project_feedback"])}

## Education & Certification Feedback

{markdown_list(result["education_feedback"])}

## Achievement Feedback

{markdown_list(result["achievement_feedback"])}

## Grammar & Language

{markdown_list(result["grammar_feedback"])}

## Bullet Point Improvements

{markdown_list(result["bullet_improvements"])}

## Recommendations

{markdown_list(result["recommendations"])}

## Final Checklist

{markdown_list(result["final_checklist"])}

---
Generated by ATS Resume Analyzer AI.
"""


# -----------------------------
# Streamlit UI
# -----------------------------

st.title("📄 ATS Resume Analyzer AI")
st.write(
    "Upload a resume and optionally provide a target job description to "
    "receive an evidence-based ATS compatibility analysis."
)

st.info(
    "The score estimates resume-to-job compatibility. It does not predict "
    "whether an employer will interview or hire a candidate."
)

with st.sidebar:
    st.header("⚙️ Settings")

    job_title = st.text_input(
        "Target Job Title",
        placeholder="e.g. Python Developer",
    )

    years_experience = st.text_input(
        "Years of Experience",
        placeholder="e.g. 2 years",
    )

    model_override = st.text_input(
        "Groq Model",
        value=DEFAULT_MODEL,
        help="Use a model currently available in your Groq account.",
    )

    st.markdown(
        """
**Privacy note:** Resume text is sent to the configured AI service for analysis.
Do not upload information you are not comfortable sending to that service.
"""
    )

col1, col2 = st.columns(2)

with col1:
    resume_file = st.file_uploader(
        "📎 Upload Resume",
        type=["pdf", "docx", "txt"],
        help="Supported formats: PDF, DOCX, TXT.",
    )

with col2:
    job_file = st.file_uploader(
        "📎 Optional Job Description File",
        type=["pdf", "docx", "txt"],
    )

job_description = st.text_area(
    "📋 Paste Job Description",
    height=220,
    placeholder="Paste the complete job description here...",
)

analyze_button = st.button(
    "🔍 Analyze Resume",
    type="primary",
    use_container_width=True,
)

if analyze_button:
    if not resume_file:
        st.error("Please upload a resume.")
        st.stop()

    try:
        with st.spinner("Extracting resume text..."):
            resume_text = extract_uploaded_file(resume_file)

        uploaded_job_text = ""

        if job_file:
            with st.spinner("Reading job description file..."):
                uploaded_job_text = extract_uploaded_file(job_file)

        combined_job_description = (job_description or "").strip()

        if uploaded_job_text:
            if combined_job_description:
                combined_job_description += "\n\n" + uploaded_job_text
            else:
                combined_job_description = uploaded_job_text

        with st.spinner("Analyzing resume with AI..."):
            result = analyze_with_groq(
                resume_text=resume_text,
                job_description=combined_job_description,
                job_title=job_title,
                years_experience=years_experience,
                model_override=model_override,
            )

        st.session_state["analysis_result"] = result
        st.session_state["resume_filename"] = resume_file.name

    except Exception as exc:
        st.error(str(exc))
        st.stop()

result = st.session_state.get("analysis_result")

if result:
    st.divider()

    c1, c2, c3 = st.columns(3)

    with c1:
        st.metric("ATS Score", f'{result["overall_score"]}/100')

    with c2:
        st.metric("Status", result["status"])

    with c3:
        keyword_count = len(result["matched_keywords"])
        st.metric("Matched Keywords", keyword_count)

    st.subheader("📝 Overall Summary")
    st.write(result["summary"])

    tabs = st.tabs(
        [
            "Overview",
            "Keywords",
            "Resume Content",
            "Formatting",
            "Improvements",
            "Report",
        ]
    )

    with tabs[0]:
        st.subheader("Category Scores")

        score_rows = []
        for key, maximum in SCORE_LIMITS.items():
            score_rows.append(
                {
                    "Category": key.replace("_", " ").title(),
                    "Score": result["category_scores"][key],
                    "Maximum": maximum,
                }
            )

        st.dataframe(
            score_rows,
            use_container_width=True,
            hide_index=True,
        )

        left, right = st.columns(2)

        with left:
            st.markdown("### ✅ Strengths")
            for item in result["strengths"]:
                st.markdown(f"- {item}")

        with right:
            st.markdown("### ⚠️ Weaknesses")
            for item in result["weaknesses"]:
                st.markdown(f"- {item}")

    with tabs[1]:
        st.subheader("🔑 Keyword Analysis")

        left, right = st.columns(2)

        with left:
            st.markdown("### Matched Keywords")
            for item in result["matched_keywords"]:
                st.markdown(f"- {item}")

            st.markdown("### Technical Skills Found")
            for item in result["technical_skills_found"]:
                st.markdown(f"- {item}")

        with right:
            st.markdown("### Missing Keywords")
            for item in result["missing_keywords"]:
                st.markdown(f"- {item}")

            st.markdown("### Technical Skills Missing")
            for item in result["technical_skills_missing"]:
                st.markdown(f"- {item}")

        st.markdown("### Unverified Keywords")
        for item in result["unverified_keywords"]:
            st.markdown(f"- {item}")

        st.markdown("### Soft Skills Found")
        for item in result["soft_skills_found"]:
            st.markdown(f"- {item}")

        st.markdown("### Soft Skills Missing")
        for item in result["soft_skills_missing"]:
            st.markdown(f"- {item}")

    with tabs[2]:
        st.subheader("Professional Summary")
        st.write(result["summary_feedback"])

        if result["improved_summary"]:
            st.markdown("### Suggested Summary")
            st.write(result["improved_summary"])

        st.subheader("Work Experience")
        for item in result["experience_feedback"]:
            st.markdown(f"- {item}")

        st.subheader("Projects")
        for item in result["project_feedback"]:
            st.markdown(f"- {item}")

        st.subheader("Education & Certifications")
        for item in result["education_feedback"]:
            st.markdown(f"- {item}")

        st.subheader("Achievements")
        for item in result["achievement_feedback"]:
            st.markdown(f"- {item}")

        st.subheader("Grammar & Language")
        for item in result["grammar_feedback"]:
            st.markdown(f"- {item}")

    with tabs[3]:
        st.subheader("ATS Formatting Issues")

        if result["formatting_issues"]:
            for item in result["formatting_issues"]:
                st.warning(item)
        else:
            st.success("No major formatting issues were identified from extracted text.")

        st.caption(
            "Formatting analysis is limited because text extraction cannot "
            "perfectly reproduce the visual layout of every PDF/DOCX."
        )

    with tabs[4]:
        st.subheader("Bullet Point Improvements")
        for item in result["bullet_improvements"]:
            st.markdown(f"- {item}")

        st.subheader("Prioritized Recommendations")
        for item in result["recommendations"]:
            st.markdown(f"- {item}")

        st.subheader("Final Checklist")
        for item in result["final_checklist"]:
            st.markdown(f"- {item}")

    with tabs[5]:
        filename = st.session_state.get(
            "resume_filename",
            "resume",
        )

        report = make_report(result, filename)

        st.download_button(
            "⬇️ Download Markdown Report",
            data=report,
            file_name="ATS_Resume_Analysis.md",
            mime="text/markdown",
            use_container_width=True,
        )

        st.download_button(
            "⬇️ Download JSON Analysis",
            data=json.dumps(result, indent=2),
            file_name="ATS_Resume_Analysis.json",
            mime="application/json",
            use_container_width=True,
        )

        with st.expander("View Raw JSON"):
            st.json(result)
